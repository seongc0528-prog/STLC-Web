"""주일예배 PPT에서 찬양 부분만 떼어 내 찬양 콘티 파일과 곡 목록 초안을 만든다.

예배 PPT(`주님의 교회 YYYY.MM.DD.pptx`)는 20MB가 넘고 설교·성경봉독까지 다 들어
있어서 그대로 올릴 수 없다. 이 스크립트는 찬양 슬라이드의 가사판 이미지만 모아
PPTX를 새로 엮고(배경 사진을 떼어 내므로 20MB대가 1~2MB로 줄어든다), 곡 수를 세어
JSON 초안을 남긴다. 2026년 10월까지 올라간 40개 콘티와 같은 모양이다.

슬라이드 구조(2025~2026년 PPT 기준):
  · 1번 = 날짜·설교 제목이 적힌 표지 → 찬양은 2번부터 시작한다.
  · 찬양이 끝나면 "예배로의 부름"이 적힌 슬라이드가 나온다 → 그 앞까지가 찬양이다.
  · 한 곡은 같은 배경 이미지를 쓰는 연속된 슬라이드다. 곡이 바뀌면 배경이 통째로
    바뀌므로, 앞 슬라이드와 겹치는 이미지가 하나도 없는 지점이 곡의 경계다.
  · 슬라이드마다 그 슬라이드에만 쓰인 이미지가 가사판이고, 여러 슬라이드가 함께
    쓰는 이미지는 배경·장식이다 — 가사판만 가져온다.

곡 제목은 가사가 이미지라서 글자로 읽을 수 없다. 그래서 곡마다 앞쪽 가사 이미지를
따로 저장해 두고(제목이 이미지 왼쪽 위에 "♬ 제목" 으로 적혀 있다), JSON의 title 은
비워 둔다. 사람이(또는 이미지를 읽을 수 있는 도구가) 채워 넣은 뒤
upload-praise-conti.mjs 로 올린다.

준비: pip install python-pptx   (이 스크립트에서만 쓰는 도구라 package.json 과는 무관)

실행:
  python scripts/extract-praise-conti.py --out <폴더>            # 아직 안 올린 주일 전부
  python scripts/extract-praise-conti.py --out <폴더> 2026.10.25  # 특정 주일만
"""

import argparse
import json
import os
import re
import sys
import zipfile

from pptx import Presentation

PPT_DIR = r"C:\Users\user\Desktop\시드니 주님의 교회\주일예배 PPT 2026"
# 파일 이름은 "주님의 교회 2026.10.25.pptx", 가끔 "..._v2.pptx" 처럼 뒤가 붙는다.
NAME_RE = re.compile(r"(\d{4})\.(\d{2})\.(\d{2})")
RELS_IMAGE_RE = re.compile(r'Target="\.\./media/([^"]+)"')
TEXT_RE = re.compile(r"<a:t>(.*?)</a:t>", re.S)
# 찬양이 끝나고 예배가 시작되는 지점. 교회 PPT가 매주 쓰는 문구다.
PRAISE_END_MARKERS = ("예배로의 부름", "예배의 부름")
FIRST_PRAISE_SLIDE = 2


def slide_text(z, index):
    xml = z.read("ppt/slides/slide%d.xml" % index).decode("utf8")
    return [t.strip() for t in TEXT_RE.findall(xml) if t.strip()]


def slide_images(z, index):
    rels = z.read("ppt/slides/_rels/slide%d.xml.rels" % index).decode("utf8")
    return RELS_IMAGE_RE.findall(rels)


def slide_count(z):
    return len([n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)])


def find_praise_range(z):
    """찬양 슬라이드의 (첫 번호, 끝 번호). 끝을 못 찾으면 None."""
    total = slide_count(z)
    for i in range(FIRST_PRAISE_SLIDE, total + 1):
        text = " ".join(slide_text(z, i))
        if any(marker in text for marker in PRAISE_END_MARKERS):
            return FIRST_PRAISE_SLIDE, i - 1
    return None


def split_songs(z, start, end):
    """배경이 통째로 바뀌는 지점을 곡 경계로 보고 슬라이드를 곡별로 묶는다."""
    songs = []
    previous = None
    for i in range(start, end + 1):
        images = set(slide_images(z, i))
        if songs and previous and images & previous:
            songs[-1].append(i)
        else:
            songs.append([i])
        previous = images
    return songs


def lyric_images(z, start, end):
    """찬양 슬라이드마다 가사판 이미지 하나씩. 못 고른 슬라이드는 None."""
    images_by_slide = {i: slide_images(z, i) for i in range(start, end + 1)}
    # 배경·장식은 여러 슬라이드가 함께 쓰고, 가사판은 그 슬라이드(와 반복 구간)만 쓴다.
    uses = {}
    for images in images_by_slide.values():
        for name in set(images):
            uses[name] = uses.get(name, 0) + 1

    picked = []
    for i in range(start, end + 1):
        candidates = [n for n in images_by_slide[i] if n.lower().endswith((".png", ".jpg", ".jpeg"))]
        if not candidates:
            picked.append(None)
            continue
        # 가장 적게 쓰인 것 → 같은 횟수면 용량이 작은 것(가사판은 글자뿐이라 가볍다).
        picked.append(min(candidates, key=lambda n: (uses[n], z.getinfo("ppt/media/" + n).file_size)))
    return picked


def write_conti(z, dest, names, out_dir):
    """가사판 이미지만 한 장에 하나씩 얹은 PPTX를 새로 만든다."""
    prs = Presentation()
    blank = prs.slide_layouts[6]
    scratch = os.path.join(out_dir, ".media")
    os.makedirs(scratch, exist_ok=True)
    for name in names:
        slide = prs.slides.add_slide(blank)
        if not name:
            continue
        path = os.path.join(scratch, name)
        if not os.path.exists(path):
            with open(path, "wb") as f:
                f.write(z.read("ppt/media/" + name))
        # 슬라이드 폭에 맞춰 위에서부터 깔면 가사판 비율(16:9)이 그대로 산다.
        slide.shapes.add_picture(path, 0, 0, width=prs.slide_width)
    prs.save(dest)


def process(path, out_dir):
    match = NAME_RE.search(os.path.basename(path))
    if not match:
        print("  날짜를 못 읽어 건너뜀: %s" % os.path.basename(path))
        return None
    year, month, day = match.groups()
    service_date = "%s-%s-%s" % (year, month, day)
    tag = (year + month + day)[2:]

    z = zipfile.ZipFile(path)
    span = find_praise_range(z)
    if not span:
        print('  "예배로의 부름" 슬라이드를 못 찾아 건너뜀: %s' % os.path.basename(path))
        return None
    start, end = span
    if end < start:
        # 3·1절처럼 특별 순서로 시작해 경배와 찬양이 아예 없는 주일이 있다.
        print("  찬양 순서가 없어 건너뜀: %s" % os.path.basename(path))
        return None
    songs = split_songs(z, start, end)
    picked = lyric_images(z, start, end)

    conti_path = os.path.join(out_dir, "%s.pptx" % tag)
    write_conti(z, conti_path, picked, out_dir)
    missing = [start + i for i, name in enumerate(picked) if name is None]
    if missing:
        print("  주의: 가사판을 못 찾은 슬라이드 %s — 콘티에서 빈 장으로 들어갑니다." % missing)

    entries = []
    for number, slides in enumerate(songs, 1):
        # 제목은 보통 첫 장 왼쪽 위에 있지만 간주·반복 가사판이 먼저 오는 주도 있어
        # 앞 세 장을 같이 꺼내 둔다.
        written = []
        for order, slide in enumerate(slides[:3], 1):
            image = picked[slide - start]
            if not image:
                continue
            name = "%s-song%d-%d%s" % (tag, number, order, os.path.splitext(image)[1])
            with open(os.path.join(out_dir, name), "wb") as f:
                f.write(z.read("ppt/media/" + image))
            written.append(name)
        entries.append(
            {
                "title": None,
                "slides": "%d-%d" % (slides[0], slides[-1]),
                "titleImages": written,
                # 슬라이드에 글자로 박힌 가사 — 제목을 가늠하는 힌트로만 쓴다.
                "slideText": " ".join(slide_text(z, slides[0]))[:120],
            }
        )

    meta = {
        "serviceDate": service_date,
        "title": "%s 찬양 콘티" % tag,
        "source": os.path.basename(path),
        "praiseSlides": "%d-%d" % (start, end),
        "file": os.path.basename(conti_path),
        "songs": entries,
    }
    meta_path = os.path.join(out_dir, "%s.json" % tag)
    with open(meta_path, "w", encoding="utf8") as f:
        json.dump(meta, f, ensure_ascii=False, indent=2)

    size = os.path.getsize(conti_path) / 1024 / 1024
    print(
        "  %s  슬라이드 %d-%d  곡 %d개  %.1fMB → %s"
        % (service_date, start, end, len(songs), size, os.path.basename(conti_path))
    )
    return meta


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("dates", nargs="*", help="처리할 주일 (예: 2026.10.25). 비우면 폴더 전체")
    parser.add_argument("--ppt-dir", default=PPT_DIR)
    parser.add_argument("--out", required=True, help="콘티 PPTX·JSON·가사 이미지를 모을 폴더")
    args = parser.parse_args()

    os.makedirs(args.out, exist_ok=True)
    files = sorted(f for f in os.listdir(args.ppt_dir) if f.lower().endswith(".pptx") and NAME_RE.search(f))
    if args.dates:
        files = [f for f in files if any(d in f for d in args.dates)]
    if not files:
        print("처리할 파일이 없습니다.")
        return 1

    print("예배 PPT %d개를 처리합니다." % len(files))
    for name in files:
        process(os.path.join(args.ppt_dir, name), args.out)
    print("\n곡 제목은 비어 있습니다. <날짜>-song<N>-<순번>.png 의 왼쪽 위 제목을 읽어")
    print("<날짜>.json 의 title 을 채운 뒤 upload-praise-conti.mjs 로 올리세요.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
